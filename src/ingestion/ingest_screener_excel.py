"""
Phase 1c: parse manually-downloaded Screener.in "Export to Excel" files and
compute ROCE, ROE, EPS, and EBITDA (none of which yfinance reliably provides
for Indian tickers).

IMPORTANT CONTEXT — read before trusting this blindly:
Screener provides no API. Earlier testing found what looked like a data bug
in Reliance's export (the annual-labeled P&L section appearing to contain
quarterly figures) — that turned out to be a parsing bug in THIS script (a
label-lookup collision, since several row labels like "Report Date", "Sales",
"Net profit", "Interest" appear twice in the sheet: once in the annual P&L
block, once again in the Quarters block). Once fixed, Reliance's annual data
matched the existing DB exactly. The lesson that stuck: never trust either
source blindly — always cross-check. So this script still cross-validates
every company's annual net profit against what's already in
`financial_statements` (from yfinance) for the SAME fiscal period, and only
trusts the annual block directly if they roughly agree. If they don't
(or there's nothing to compare against), it falls back to summing the last
4 quarters (TTM) instead, which was independently verified reliable.

Balance Sheet data (Equity, Reserves, Borrowings, No. of Equity Shares) is
used directly either way — it's a point-in-time snapshot with no
summing-quarters step to go wrong, and it's also NOT one of the duplicated
labels (unlike Sales/Net profit/PBT/Interest), so there's no collision risk.

NOT included: intrinsic value / DCF. That requires real modeling decisions
(growth assumptions, discount rate, terminal value) and belongs in its own
dedicated Phase 2 fundamentals-engine feature, not bolted onto a parser.

Manual step required: log into screener.in (free account), open each
watchlist company's CONSOLIDATED page, click "Export to Excel", and save
the file into data/screener_exports/ named as SYMBOL.xlsx (e.g. RELIANCE.xlsx,
HDFCBANK.xlsx — matching the root of your NSE ticker, no ".NS" suffix).

Run locally: python -m src.ingestion.ingest_screener_excel [SYMBOL.xlsx]
             python -m src.ingestion.ingest_screener_excel        (does all files in the folder)
"""

import sys
from pathlib import Path

import openpyxl
from sqlalchemy import select

from src.models.db import SessionLocal, init_db
from src.models.schema import Security, Fundamentals, FinancialStatement

PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXPORTS_DIR = PROJECT_ROOT / "data" / "screener_exports"

# If Screener's annual-labeled net profit differs from the DB's by more than
# this ratio, treat the annual block as unreliable and fall back to TTM.
SANITY_CHECK_TOLERANCE = 0.35  # allow +/-35%


def _row_dict(all_rows, start_row: int, end_row: int):
    """Map row label (column A) -> row index, restricted to [start_row, end_row]
    inclusive (1-indexed). Restricting range matters for ambiguous labels —
    see module docstring for why."""
    labels = {}
    for i in range(start_row, min(end_row, len(all_rows)) + 1):
        label = all_rows[i - 1][0]
        if label is not None:
            labels[label] = i
    return labels


def parse_screener_file(path: Path):
    """Returns (annual, quarters) dicts, each label -> tuple of values."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Data Sheet"]
    all_rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

    def get(row_num):
        return all_rows[row_num - 1][1:]

    quarters_marker = None
    for i, row in enumerate(all_rows, start=1):
        if row[0] == "Quarters":
            quarters_marker = i
            break
    if quarters_marker is None:
        raise ValueError(f"{path.name}: could not find 'Quarters' section marker")

    # AMBIGUOUS labels (Report Date, Sales, Net profit, PBT, Interest,
    # Depreciation) appear again in the Quarters block further down, so this
    # lookup is bounded to strictly before the Quarters marker to force
    # matching the first (annual-labeled) occurrence only.
    pl_labels = _row_dict(all_rows, start_row=1, end_row=quarters_marker - 1)

    # UNIQUE labels (Equity Share Capital, Reserves, Borrowings, No. of
    # Equity Shares) have no Quarters-block duplicate and physically appear
    # AFTER the Quarters block, so a whole-sheet lookup is safe and necessary.
    bs_labels = _row_dict(all_rows, start_row=1, end_row=len(all_rows))

    annual = {
        "report_date": get(pl_labels["Report Date"]),
        "sales": get(pl_labels["Sales"]),
        "net_profit": get(pl_labels["Net profit"]),
        "pbt": get(pl_labels["Profit before tax"]),
        "interest": get(pl_labels["Interest"]),
        "depreciation": get(pl_labels["Depreciation"]),
        "equity": get(bs_labels["Equity Share Capital"]),
        "reserves": get(bs_labels["Reserves"]),
        "borrowings": get(bs_labels["Borrowings"]),
        "num_shares": get(bs_labels["No. of Equity Shares"]),
    }

    # Quarters block: bounded to just after the marker.
    quarters_labels = _row_dict(all_rows, start_row=quarters_marker, end_row=quarters_marker + 14)
    quarters = {
        "net_profit": get(quarters_labels["Net profit"]),
        "pbt": get(quarters_labels["Profit before tax"]),
        "interest": get(quarters_labels["Interest"]),
        "depreciation": get(quarters_labels["Depreciation"]),
    }

    return annual, quarters


def get_security(session, symbol: str):
    return session.execute(
        select(Security).where(Security.symbol == symbol)
    ).scalar_one_or_none()


def matching_db_net_profit(session, security_id: int, target_period_end):
    """Net profit from financial_statements for the SAME period Screener's
    export ends on — not just whichever row happens to be latest in the DB.
    Falls back to the nearest available annual period within 60 days if an
    exact match isn't found, and flags when that fallback happened."""
    stmt = (
        select(FinancialStatement)
        .where(
            FinancialStatement.security_id == security_id,
            FinancialStatement.period_type == "annual",
            FinancialStatement.currency == "INR",
        )
        .order_by(FinancialStatement.period_end.desc())
    )
    candidates = session.execute(stmt).scalars().all()
    if not candidates:
        return None, None, False

    for row in candidates:
        if row.period_end == target_period_end:
            return row.net_profit, row.period_end, True

    nearest = min(candidates, key=lambda r: abs((r.period_end - target_period_end).days))
    if abs((nearest.period_end - target_period_end).days) <= 60:
        return nearest.net_profit, nearest.period_end, False
    return None, None, False


def compute_ratios(annual: dict, quarters: dict, db_net_profit_inr):
    """Returns a dict of computed ratios plus which method was used
    ('annual-direct' or 'ttm-fallback')."""
    screener_latest_net_profit_inr = annual["net_profit"][-1] * 1e7

    if db_net_profit_inr:
        ratio = screener_latest_net_profit_inr / db_net_profit_inr
        use_annual_block = (1 - SANITY_CHECK_TOLERANCE <= ratio <= 1 + SANITY_CHECK_TOLERANCE)
    else:
        # No DB data to verify against — default to the safer TTM method
        # rather than trusting unverified annual data.
        use_annual_block = False

    equity, reserves, borrowings = annual["equity"], annual["reserves"], annual["borrowings"]
    cap_emp_latest = equity[-1] + reserves[-1] + borrowings[-1]
    cap_emp_prev = equity[-2] + reserves[-2] + borrowings[-2]
    num_shares_latest = annual["num_shares"][-1]

    if use_annual_block:
        pbt = annual["pbt"][-1]
        interest = annual["interest"][-1]
        depreciation = annual["depreciation"][-1]
        net_profit = annual["net_profit"][-1]
        method = "annual-direct"
    else:
        pbt = sum(quarters["pbt"][-4:])
        interest = sum(quarters["interest"][-4:])
        depreciation = sum(quarters["depreciation"][-4:])
        net_profit = sum(quarters["net_profit"][-4:])
        method = "ttm-fallback"

    ebit = pbt + interest
    ebitda_cr = ebit + depreciation  # Rs. Crore
    roce = ebit / ((cap_emp_latest + cap_emp_prev) / 2) if (cap_emp_latest + cap_emp_prev) else None
    roe = net_profit / (equity[-1] + reserves[-1]) if (equity[-1] + reserves[-1]) else None
    eps = (net_profit * 1e7) / num_shares_latest if num_shares_latest else None  # Rs. per share

    return {
        "method": method,
        "roce": roce,
        "roe": roe,
        "eps": eps,
        "ebitda": ebitda_cr * 1e7,  # convert Rs. Crore -> raw INR, to match financial_statements convention
    }


def update_fundamentals(session, security_id: int, ratios: dict):
    """Updates ROCE, EPS, and EBITDA on the most recent fundamentals row.
    ROE is intentionally left as-is (yfinance's value) unless it's missing —
    Screener's ROE is computed too, but only fills a NULL, since yfinance's
    ROE has already been working fine and there's no evidence it needs
    fixing, unlike ROCE which yfinance never populated at all."""
    stmt = (
        select(Fundamentals)
        .where(Fundamentals.security_id == security_id)
        .order_by(Fundamentals.as_of_date.desc())
    )
    row = session.execute(stmt).scalars().first()
    if row is None:
        print(f"  No existing fundamentals row for security_id={security_id} — "
              f"run ingest_stock_price first. Skipping.")
        return

    row.roce = ratios["roce"]
    row.eps = ratios["eps"]
    row.ebitda = ratios["ebitda"]
    if row.roe is None:
        row.roe = ratios["roe"]
    session.commit()


def process_file(path: Path):
    symbol = path.stem.upper() + ".NS"
    print(f"\n--- {path.name} -> {symbol} ---")

    annual, quarters = parse_screener_file(path)

    init_db()
    session = SessionLocal()
    try:
        security = get_security(session, symbol)
        if security is None:
            print(f"  {symbol} not found in securities table — "
                  f"run ingest_stock_price for it first. Skipping.")
            return

        screener_latest_period = annual["report_date"][-1].date()
        db_net_profit, db_period, exact_match = matching_db_net_profit(
            session, security.id, screener_latest_period
        )

        if db_net_profit:
            match_note = ("exact period match" if exact_match
                           else f"nearest available, not exact — DB has {db_period}, "
                                f"export has {screener_latest_period}")
            print(f"  DB net profit ({match_note}): {db_net_profit:,.0f} INR")
        else:
            print(f"  No DB net profit found within range for cross-check — "
                  f"defaulting to TTM fallback for safety")

        ratios = compute_ratios(annual, quarters, db_net_profit)

        print(f"  Method used: {ratios['method']}")
        print(f"  ROCE:   {ratios['roce']:.2%}" if ratios['roce'] is not None else "  ROCE: N/A")
        print(f"  ROE:    {ratios['roe']:.2%}" if ratios['roe'] is not None else "  ROE: N/A")
        print(f"  EPS:    Rs {ratios['eps']:.2f}" if ratios['eps'] is not None else "  EPS: N/A")
        print(f"  EBITDA: Rs {ratios['ebitda']:,.0f} ({ratios['ebitda']/1e7:,.0f} cr)")

        update_fundamentals(session, security.id, ratios)
        print(f"  Updated fundamentals for {symbol}")
    finally:
        session.close()


def main():
    if len(sys.argv) > 1:
        files = [EXPORTS_DIR / sys.argv[1]]
    else:
        files = sorted(EXPORTS_DIR.glob("*.xlsx"))

    if not files:
        print(f"No .xlsx files found in {EXPORTS_DIR}. "
              f"Export from screener.in and save files there first.")
        return

    for f in files:
        if not f.exists():
            print(f"File not found: {f}")
            continue
        process_file(f)


if __name__ == "__main__":
    main()